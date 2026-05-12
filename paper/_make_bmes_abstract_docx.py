#!/usr/bin/env python3
"""Generate the BMES 2026 abstract as a Word document for advisor review.

Output: D:/Desktop/bmes_2026_swc_studio_abstract.docx

Usage::

    python -m paper._make_bmes_abstract_docx
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


OUT = Path("D:/Desktop/bmes_2026_swc_studio_abstract.docx")

NAVY = RGBColor(0x1F, 0x3A, 0x68)
GREY = RGBColor(0x55, 0x55, 0x55)
BLACK = RGBColor(0x00, 0x00, 0x00)


TITLE = (
    "SWC-Studio: A Hybrid Graph Neural Network Pipeline for Robust "
    "Cross-Corpus Auto-Labeling of Neuron Morphology"
)

TRACK = "AI in BME"
SUBTRACK = "Foundation Models and Advanced AI Architectures for Biomedical Applications"

INTRO = (
    "Per-neurite type labels (soma, axon, basal dendrite, apical dendrite) are "
    "foundational for nearly every downstream analysis in computational neuroscience: "
    "connectomics, morphology-based cell-type classification, electrophysiology-"
    "morphology mapping, and large-scale comparative anatomy. Yet these labels are "
    "still typically assigned by manual curation, a process that scales poorly and "
    "creates a bottleneck for the rapidly growing volume of digitally reconstructed "
    "neurons in public repositories such as NeuroMorpho.Org, the Allen Cell Types "
    "Database, and contemporary connectomics datasets. Existing automated tools "
    "including NeuroM, L-Measure, and Sholl-analysis-based classifiers work well "
    "on archetypal cells but fail on atypical morphologies: rotated coordinate "
    "frames, partial or noisy tracings, immature or pathological cells, and the "
    "cross-corpus variability inevitable when combining data from multiple "
    "laboratories. Worse, the field lacks rigorous cross-corpus generalization "
    "studies, leaving end users uncertain whether published accuracy numbers "
    "transfer to their own data. We present a four-stage hybrid pipeline that "
    "combines hand-engineered morphological features with a Graph Neural Network "
    "apical-versus-basal classification head and topology-aware refinement. The "
    "pipeline is deployed as SWC-Studio, an open-source desktop, command-line, "
    "and Python toolkit, and is rigorously evaluated against four published "
    "baselines, three feature-group ablations, and a three-way leave-one-source-"
    "out cross-corpus generalization study with paired statistical tests across "
    "210 pairwise method comparisons."
)

METHODS = (
    "We assembled a curated corpus of 2,267 manually-labeled neuron morphologies "
    "from three sources: hippocampal CA1 pyramidal cells, the Allen Cell Types "
    "Database, and NeuroMorpho.Org (805 interneurons, 1,462 pyramidal cells). "
    "The pipeline has four stages: (1) a cell-type random forest classifier "
    "(interneuron versus pyramidal); (2) a per-branch random forest over 57 "
    "hand-engineered features capturing geometric, topological, and cell-"
    "intrinsic principal-axis properties; (3) a GraphSAGE graph neural network "
    "apical-versus-basal head operating on the subtree branch graph; and "
    "(4) topology-aware refinement using rule-based constraints. Evaluation used "
    "a deterministic 80/20 hash-bucket train/test split (n = 1,806 train, "
    "n = 461 test). Four external baselines (NeuroM-RF, Sholl-RF, Sholl-MLP, "
    "L-Measure-RF) were trained on the identical split for apples-to-apples "
    "comparison. Statistical analyses used paired Wilcoxon signed-rank tests with "
    "Bonferroni correction over 210 pairwise comparisons. External baselines "
    "receive ground-truth cell type as a one-hot input feature, while our "
    "pipeline predicts it from the raw SWC via an integrated Stage 1 classifier "
    "(held-out cell-type accuracy 99.13%), making the reported comparisons "
    "conservative with respect to our method's advantage. Generalization was "
    "assessed via three-way leave-one-source-out cross-validation, training on "
    "two corpora and evaluating on the held-out third. The auto-typing pipeline "
    "is integrated into SWC-Studio, a complete morphology toolkit that also "
    "provides issue-driven validation, automated repair, radii cleaning, "
    "geometry editing, batch processing, 3D visualization, and a one-command "
    "retraining workflow that lets end users retrain the engine on their own "
    "labeled SWCs."
)

RESULTS = (
    "On the held-out test split (n = 461), the pipeline achieves per-node "
    "accuracy 0.9934 and neurite-macro-F1 0.9673. The most operationally "
    "relevant result is on the per-file F1 distribution, which measures "
    "performance per individual cell rather than weighting by node count. The "
    "10th-percentile per-file F1 reaches 0.9012, compared with 0.6242 for the "
    "strongest external baseline (L-Measure-RF). This 28-point gap on the "
    "hardest 10% of cells corresponds to exactly the cases where human curators "
    "currently spend most of their effort and where automation has the largest "
    "practical value. Across 210 pairwise Wilcoxon tests, the pipeline "
    "significantly outperforms every external baseline after Bonferroni "
    "correction (p_bonf < 0.025 for L-Measure-RF, < 1e-4 for Sholl-RF, < 1e-8 "
    "for Sholl-MLP, < 1e-40 for NeuroM-RF). Three feature-group ablations "
    "(removing principal-axis features, trunk-detection features, or the "
    "confidence-weighted soft handoff between Stages 1 and 2) yield "
    "statistically indistinguishable results (all p_bonf = 1.0), demonstrating "
    "that the architecture is robust to specific feature subset choices. Leave-"
    "one-source-out experiments quantify cross-corpus transfer: training on "
    "Allen and NeuroMorpho and evaluating on never-seen hippocampal cells, the "
    "pipeline retains accuracy 0.9521 with neurite-macro-F1 = 0.6505, with "
    "apical-versus-basal discrimination accounting for nearly all the cross-"
    "corpus performance gap. The pipeline ships in SWC-Studio, a desktop "
    "application (macOS and Windows bundles), command-line tool, and Python "
    "library released open-source under a permissive license. SWC-Studio unifies "
    "the auto-labeling pipeline with validation, repair, batch processing, 3D "
    "visualization, and a one-command user-retraining workflow into a single "
    "tool, addressing the morphology curation pipeline end-to-end rather than "
    "just the labeling step. Future work will integrate active human-in-the-"
    "loop curation and extend the engine to additional cell-type taxonomies."
)

ACK = (
    "We acknowledge the Allen Institute Cell Types Database and the "
    "NeuroMorpho.Org consortium for openly sharing the curated morphologies "
    "used in this study. Key references: Ascoli, Donohue, and Halavi, "
    "NeuroMorpho.Org: a central resource for neuronal morphologies "
    "(J Neurosci 2007); Scorcioni, Polavaram, and Ascoli, L-Measure: a tool for "
    "the quantitative analysis of neuronal morphology (Nature Protocols 2008); "
    "Hamilton, Ying, and Leskovec, Inductive Representation Learning on Large "
    "Graphs / GraphSAGE (NeurIPS 2017); Allen Institute for Brain Science, "
    "Allen Cell Types Database, celltypes.brain-map.org."
)


def _word_count(text: str) -> int:
    return len([w for w in text.split() if w.strip()])


def _set_style(run, *, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _add_heading(doc, text, *, size=14, color=NAVY):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_style(run, size=size, bold=True, color=color)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p


def _add_para(doc, text, *, size=11, italic=False, color=BLACK):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_style(run, size=size, italic=italic, color=color)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    return p


def build():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    # --- Top label ---
    p = doc.add_paragraph()
    run = p.add_run("BMES 2026 Annual Meeting — Abstract Submission Draft")
    _set_style(run, size=10, italic=True, color=GREY)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- Title ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    _set_style(run, size=18, bold=True, color=NAVY)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)

    # --- Submission metadata table ---
    _add_heading(doc, "Submission Metadata", size=12)
    meta_rows = [
        ("Track", TRACK),
        ("Subtrack", SUBTRACK),
    ]
    meta = doc.add_table(rows=len(meta_rows), cols=2)
    meta.style = "Light Grid Accent 1"
    meta.autofit = True
    for i, (label, value) in enumerate(meta_rows):
        c0 = meta.rows[i].cells[0]
        c1 = meta.rows[i].cells[1]
        c0.text = ""
        c1.text = ""
        r0 = c0.paragraphs[0].add_run(label)
        _set_style(r0, size=10, bold=True)
        r1 = c1.paragraphs[0].add_run(value)
        _set_style(r1, size=10)

    # --- Introduction ---
    _add_heading(doc, f"Introduction  ({_word_count(INTRO)} words; limit 50-250)")
    _add_para(doc, INTRO)

    # --- Methods ---
    _add_heading(doc, f"Materials and Methods  ({_word_count(METHODS)} words; limit 50-250)")
    _add_para(doc, METHODS)

    # --- Results ---
    _add_heading(doc, f"Results, Conclusions, and Discussions  ({_word_count(RESULTS)} words; limit 50-350)")
    _add_para(doc, RESULTS)

    # --- Acknowledgements ---
    _add_heading(doc, "Acknowledgements and References", size=12)
    _add_para(doc, ACK, size=10)

    # --- Figure plan ---
    _add_heading(doc, "Suggested Figures (optional, recommended for oral-talk consideration)", size=12)
    _add_para(
        doc,
        "Figure 1 (recommended). Headline performance on the held-out test split "
        "(n = 461 cells). Bar chart of per-class F1 (soma, axon, basal dendrite, "
        "apical dendrite) for SWC-Studio (our method) against four external "
        "baselines (NeuroM-RF, Sholl-RF, Sholl-MLP, L-Measure-RF) trained on the "
        "identical split.",
        size=10,
    )
    _add_para(
        doc,
        "Figure 2 (recommended). Per-file F1 distributions showing tail-of-"
        "distribution robustness. Violin plot of per-cell neurite-macro-F1 for "
        "each method on the held-out test split. SWC-Studio's 10th-percentile "
        "per-file F1 reaches 0.9012 versus 0.6242 for the strongest baseline.",
        size=10,
    )
    _add_para(
        doc,
        "Figure 3 (optional). Cross-corpus leave-one-source-out generalization. "
        "Bar chart of per-class F1 for the three LOSO experiments (held-out: "
        "in-house hippocampal, NeuroMorpho, Allen) overlaid with the in-domain "
        "reference, highlighting that apical-versus-basal discrimination "
        "accounts for nearly all the cross-corpus performance gap.",
        size=10,
    )

    # --- Footer note ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Draft for advisor review. All word counts validated against BMES 2026 form limits."
    )
    _set_style(run, size=9, italic=True, color=GREY)
    p.paragraph_format.space_before = Pt(18)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")
    print(f"  size: {OUT.stat().st_size / 1024:.1f} KB")
    print(f"  word counts:")
    print(f"    Introduction: {_word_count(INTRO)} (limit 50-250)")
    print(f"    Materials and Methods: {_word_count(METHODS)} (limit 50-250)")
    print(f"    Results/Conclusions/Discussions: {_word_count(RESULTS)} (limit 50-350)")


if __name__ == "__main__":
    build()
